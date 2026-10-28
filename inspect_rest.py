import pptx
prs = pptx.Presentation('/home/istiaqfuad/Desktop/gigalogy_project/backend_architecture_presentation.pptx')
for i in range(4, 9):
    slide = prs.slides[i]
    print(f"\n--- Slide {i} ---")
    blocks = 0
    title = ""
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = shape.text.replace(chr(10), ' ')[:40]
            if shape.name == "TextBox 2": title = text
            if "Rectangle" in shape.name: blocks += 1
    print(f"Title: {title}")
    print(f"Blocks: {blocks}")
