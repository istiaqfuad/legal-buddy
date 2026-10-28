import pptx
prs = pptx.Presentation('/home/istiaqfuad/Desktop/gigalogy_project/backend_architecture_presentation.pptx')
slide = prs.slides[0]
print("Shapes on Slide 0:")
for shape in slide.shapes:
    print(f"Shape name: {shape.name}, type: {shape.shape_type}")
    if shape.has_text_frame:
        print(f"  Text: {shape.text[:30]}")
