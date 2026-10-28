import pptx
prs = pptx.Presentation('/home/istiaqfuad/Desktop/gigalogy_project/ai_architecture_presentation.pptx')
for i, slide in enumerate(prs.slides[:3]):
    print(f"\n--- Slide {i} ---")
    for shape in slide.shapes:
        if shape.has_text_frame:
            print(f"{shape.name} Text: {shape.text.replace(chr(10), ' ')[:40]}")
